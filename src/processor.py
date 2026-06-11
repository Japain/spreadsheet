import time
from datetime import datetime
from pathlib import Path

import openpyxl
from PySide6.QtCore import QThread, Signal

from src.config import Config, Workbook, DEFAULT_LOG_PATH
from src.log import RunRecord, RunResult, append_record
from src.validation import SUFFIX_RE as _SUFFIX_RE


def _output_path(wb: Workbook, suffix: str) -> Path:
    """Canonical output path for a workbook: {folder}/{stem}_{suffix}.xlsx."""
    return Path(wb.folder) / f"{Path(wb.filename).stem}_{suffix}.xlsx"


def check_conflicts(config: Config, selected_ids: list[str], suffix: str) -> list[str]:
    """Return output filenames that already exist on disk for the given selection."""
    if not _SUFFIX_RE.match(suffix):
        raise ValueError("Suffix must contain only alphanumeric characters, underscores, or dashes.")
    wb_map = {wb.id: wb for wb in config.workbooks}
    conflicts = []
    for wb_id in selected_ids:
        wb = wb_map.get(wb_id)
        if wb is None:
            continue
        out = _output_path(wb, suffix)
        if out.exists():
            conflicts.append(out.name)
    return conflicts


class WorkbookProcessor(QThread):
    workbook_started = Signal(str)
    workbook_finished = Signal(RunResult)
    run_complete = Signal(list)
    run_error = Signal(str)

    def configure(
        self,
        config: Config,
        input_filename: str,
        suffix: str,
        selected_ids: list[str],
    ) -> None:
        if self.isRunning():
            raise RuntimeError("Cannot configure processor while it is running.")
        if not _SUFFIX_RE.match(suffix):
            raise ValueError("Suffix must contain only alphanumeric characters, underscores, or dashes.")
        self._config = config
        self._input_filename = input_filename
        self._suffix = suffix
        self._selected_ids = selected_ids

    def run(self) -> None:
        config = self._config
        suffix = self._suffix
        input_filename = self._input_filename
        selected_ids = self._selected_ids

        master_path = Path(config.input_folder) / input_filename
        if not master_path.exists():
            self.run_error.emit(f"Master file not found: {master_path}")
            return

        try:
            master_wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
        except Exception as exc:
            self.run_error.emit(f"Could not open master file: {exc}")
            return

        wb_map = {wb.id: wb for wb in config.workbooks}
        results: list[RunResult] = []

        try:
            for wb_id in selected_ids:
                self.workbook_started.emit(wb_id)
                t0 = time.monotonic()

                wb_cfg = wb_map.get(wb_id)
                if wb_cfg is None:
                    r = RunResult(
                        workbook_id=wb_id,
                        filename=wb_id,
                        status="error",
                        message=f"Workbook '{wb_id}' not found in config",
                        output_filename=None,
                        rows_written=0,
                        duration_ms=int((time.monotonic() - t0) * 1000),
                    )
                else:
                    r = self._process_single_workbook(wb_cfg, master_wb, suffix, t0)

                self.workbook_finished.emit(r)
                results.append(r)
        finally:
            master_wb.close()

        try:
            record = RunRecord(
                timestamp=datetime.now().isoformat(timespec="seconds"),
                input_filename=input_filename,
                suffix=suffix,
                results=results,
            )
            append_record(DEFAULT_LOG_PATH, record)
        except Exception as exc:
            self.run_error.emit(f"Could not write run log: {exc}")
            return

        self.run_complete.emit(results)

    def _process_single_workbook(
        self,
        wb_cfg: Workbook,
        master_wb: openpyxl.Workbook,
        suffix: str,
        t0: float,
    ) -> RunResult:
        output_path = _output_path(wb_cfg, suffix)
        target_path = Path(wb_cfg.folder) / wb_cfg.filename

        def _error(msg: str) -> RunResult:
            return RunResult(
                workbook_id=wb_cfg.id,
                filename=wb_cfg.filename,
                status="error",
                message=msg,
                output_filename=None,
                rows_written=0,
                duration_ms=int((time.monotonic() - t0) * 1000),
            )

        if not target_path.exists():
            return _error(f"Target file not found: {wb_cfg.filename}")

        try:
            with open(target_path, "r+b"):
                pass
        except OSError:
            return _error("File could not be accessed — it may be locked or you may not have write permission")

        target_wb = None
        try:
            try:
                target_wb = openpyxl.load_workbook(target_path)
            except Exception as exc:
                return _error(f"Could not open target file: {exc}")

            any_skipped = False
            rows_written = 0

            for mapping in wb_cfg.mappings:
                if mapping.input not in master_wb.sheetnames:
                    any_skipped = True
                    continue
                if mapping.target not in target_wb.sheetnames:
                    any_skipped = True
                    continue

                master_ws = master_wb[mapping.input]
                target_ws = target_wb[mapping.target]

                # Reset XML-derived dimension cache; load all rows to get true dimensions.
                # openpyxl read-only mode reads max_column/max_row from the <dimension> tag,
                # which is often wrong or missing in files from pandas/xlsxwriter/Google Sheets.
                master_ws.reset_dimensions()
                master_rows = list(master_ws.iter_rows(values_only=True))
                master_max_row = len(master_rows)
                master_max_col = max((len(row) for row in master_rows), default=0)
                target_max_row = target_ws.max_row or 0

                for row in range(1, target_max_row + 1):
                    for col in range(1, master_max_col + 1):
                        target_ws.cell(row=row, column=col).value = None

                for r_i, row_vals in enumerate(master_rows, start=1):
                    for c_i, value in enumerate(row_vals, start=1):
                        target_ws.cell(row=r_i, column=c_i).value = value

                rows_written += master_max_row

            try:
                target_wb.save(output_path)
            except (PermissionError, OSError) as exc:
                return _error(f"Could not save output: {exc}")

            status = "skipped" if any_skipped else "success"
            return RunResult(
                workbook_id=wb_cfg.id,
                filename=wb_cfg.filename,
                status=status,
                message="",
                output_filename=output_path.name,
                rows_written=rows_written,
                duration_ms=int((time.monotonic() - t0) * 1000),
            )
        finally:
            if target_wb is not None:
                target_wb.close()
