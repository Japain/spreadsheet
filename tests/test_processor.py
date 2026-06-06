"""Processor tests — uses real .xlsx files via openpyxl; no file I/O mocking."""
import os
import re
import sys
import zipfile
import pytest
import openpyxl
from pathlib import Path

import src.processor as _proc_module
from src.config import Config, Workbook, TabMapping
from src.log import RunResult
from src.processor import check_conflicts, WorkbookProcessor


@pytest.fixture(autouse=True)
def _redirect_log(tmp_path, monkeypatch):
    monkeypatch.setattr(_proc_module, "DEFAULT_LOG_PATH", tmp_path / "test_runs.log")


# ── helpers ──────────────────────────────────────────────────────────────────

def _wb(path: Path, sheets: dict) -> None:
    """Create a workbook. sheets: {name: [[row], ...]} (first sheet replaces default)."""
    book = openpyxl.Workbook()
    first = next(iter(sheets))
    book.active.title = first
    for name, rows in sheets.items():
        ws = book[name] if name in book.sheetnames else book.create_sheet(name)
        for row in rows:
            ws.append(row)
    book.save(path)


def _cfg(input_folder: str, workbooks: list[Workbook]) -> Config:
    return Config(input_folder=input_folder, workbooks=workbooks)


def _workbook(id: str, filename: str, folder: str, mappings: list[tuple]) -> Workbook:
    return Workbook(
        id=id,
        filename=filename,
        folder=folder,
        mappings=[TabMapping(input=i, target=t) for i, t in mappings],
    )


class _Cap:
    """Connects to WorkbookProcessor signals and captures emitted values."""
    def __init__(self, p: WorkbookProcessor):
        self.started: list[str] = []
        self.finished: list[RunResult] = []
        self.complete: list[list] = []
        self.errors: list[str] = []
        p.workbook_started.connect(self.started.append)
        p.workbook_finished.connect(self.finished.append)
        p.run_complete.connect(self.complete.append)
        p.run_error.connect(self.errors.append)


# ── check_conflicts ──────────────────────────────────────────────────────────

def test_check_conflicts_empty_when_no_output_file_exists(tmp_path):
    folder = tmp_path / "out"
    folder.mkdir()
    _wb(folder / "report.xlsx", {"Sheet1": []})
    config = _cfg("", [_workbook("wb1", "report.xlsx", str(folder), [("Sheet1", "Sheet1")])])

    assert check_conflicts(config, ["wb1"], "Q1") == []


def test_check_conflicts_returns_conflicting_filenames(tmp_path):
    folder = tmp_path / "out"
    folder.mkdir()
    _wb(folder / "report.xlsx", {"Sheet1": []})
    (folder / "report_Q1.xlsx").touch()
    config = _cfg("", [_workbook("wb1", "report.xlsx", str(folder), [("Sheet1", "Sheet1")])])

    assert check_conflicts(config, ["wb1"], "Q1") == ["report_Q1.xlsx"]


def test_check_conflicts_invalid_suffix_raises_value_error():
    with pytest.raises(ValueError, match="Suffix"):
        check_conflicts(_cfg("", []), [], "../../bad")


# ── WorkbookProcessor.run() ───────────────────────────────────────────────────

def test_successful_run_writes_values_leaves_source_unchanged(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["A", "B"], ["1", "2"]]})
    _wb(out_dir / "report.xlsx", {"Sheet1": []})

    config = _cfg(str(master_dir), [_workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    assert len(cap.finished) == 1
    assert cap.finished[0].status == "success"
    assert cap.complete == [[cap.finished[0]]]

    # Output file has correct values
    out = openpyxl.load_workbook(out_dir / "report_Q1.xlsx")
    ws = out["Sheet1"]
    assert ws.cell(1, 1).value == "A"
    assert ws.cell(1, 2).value == "B"
    assert ws.cell(2, 1).value == "1"
    assert ws.cell(2, 2).value == "2"

    # Source (original target) is unchanged — no data written into it
    src = openpyxl.load_workbook(out_dir / "report.xlsx")
    assert all(cell.value is None for row in src["Sheet1"].iter_rows() for cell in row)
    out.close()
    src.close()


def test_master_not_found_emits_run_error_no_output_created(tmp_path, qapp):
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    _wb(out_dir / "report.xlsx", {"Sheet1": []})

    config = _cfg(str(tmp_path / "master"), [
        _workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "nonexistent.xlsx", "Q1", ["wb1"])
    p.run()

    assert len(cap.errors) == 1
    assert cap.finished == []
    assert not (out_dir / "report_Q1.xlsx").exists()


def test_target_not_found_emits_error_and_continues_other_workbooks(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]]})
    _wb(out_dir / "ok.xlsx", {"Sheet1": []})
    # "missing.xlsx" intentionally absent

    config = _cfg(str(master_dir), [
        _workbook("missing", "missing.xlsx", str(out_dir), [("Sheet1", "Sheet1")]),
        _workbook("ok", "ok.xlsx", str(out_dir), [("Sheet1", "Sheet1")]),
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["missing", "ok"])
    p.run()

    assert cap.errors == []
    assert len(cap.finished) == 2
    statuses = {r.workbook_id: r.status for r in cap.finished}
    assert statuses["missing"] == "error"
    assert statuses["ok"] == "success"
    assert (out_dir / "ok_Q1.xlsx").exists()


def test_missing_input_tab_skips_mapping_saves_with_skipped_status(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]]})
    _wb(out_dir / "report.xlsx", {"Sheet1": [], "Sheet2": []})

    config = _cfg(str(master_dir), [
        _workbook("wb1", "report.xlsx", str(out_dir), [
            ("NoSuchTab", "Sheet1"),  # missing in master
            ("Sheet1", "Sheet2"),
        ])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    assert len(cap.finished) == 1
    assert cap.finished[0].status == "skipped"
    assert (out_dir / "report_Q1.xlsx").exists()


def test_missing_target_tab_skips_mapping_saves_with_skipped_status(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]], "Sheet2": [["other"]]})
    _wb(out_dir / "report.xlsx", {"Sheet1": []})  # Sheet2 absent in target

    config = _cfg(str(master_dir), [
        _workbook("wb1", "report.xlsx", str(out_dir), [
            ("Sheet1", "Sheet1"),
            ("Sheet2", "NoSuchTarget"),  # missing in target
        ])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    assert len(cap.finished) == 1
    assert cap.finished[0].status == "skipped"
    assert (out_dir / "report_Q1.xlsx").exists()


def test_paste_zone_cleared_before_write(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    # Master has 2 data rows
    _wb(master_dir / "master.xlsx", {"Sheet1": [["new1"], ["new2"]]})
    # Target has 5 stale rows
    _wb(out_dir / "report.xlsx", {"Sheet1": [["old1"], ["old2"], ["old3"], ["old4"], ["old5"]]})

    config = _cfg(str(master_dir), [
        _workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    out = openpyxl.load_workbook(out_dir / "report_Q1.xlsx")
    ws = out["Sheet1"]
    assert ws.cell(1, 1).value == "new1"
    assert ws.cell(2, 1).value == "new2"
    # Stale rows 3–5 must be cleared
    assert ws.cell(3, 1).value is None
    assert ws.cell(4, 1).value is None
    assert ws.cell(5, 1).value is None
    out.close()


def test_values_only_write_no_formula_strings_in_output(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    # Create master with a formula cell
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    ws["B1"] = "=A1*2"  # formula — should NOT appear as-is in output
    book.save(master_dir / "master.xlsx")
    _wb(out_dir / "report.xlsx", {"Sheet1": []})

    config = _cfg(str(master_dir), [
        _workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    out = openpyxl.load_workbook(out_dir / "report_Q1.xlsx")
    for row in out["Sheet1"].iter_rows():
        for cell in row:
            if cell.value is not None:
                assert not str(cell.value).startswith("="), (
                    f"{cell.coordinate} contains formula: {cell.value!r}"
                )
    out.close()


def test_columns_right_of_max_column_untouched(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    # Master has 2 columns
    _wb(master_dir / "master.xlsx", {"Sheet1": [["m1", "m2"]]})
    # Target has data in columns 3–4 (outside master's paste zone)
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Sheet1"
    ws["A1"] = "old_a"
    ws["B1"] = "old_b"
    ws["C1"] = "keep_c"
    ws["D1"] = "keep_d"
    book.save(out_dir / "report.xlsx")

    config = _cfg(str(master_dir), [
        _workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    out = openpyxl.load_workbook(out_dir / "report_Q1.xlsx")
    ws = out["Sheet1"]
    assert ws.cell(1, 1).value == "m1"
    assert ws.cell(1, 2).value == "m2"
    assert ws.cell(1, 3).value == "keep_c"
    assert ws.cell(1, 4).value == "keep_d"
    out.close()


@pytest.mark.skipif(sys.platform == "win32", reason="chmod-based lock simulation not reliable on Windows")
def test_locked_file_emits_error_and_continues(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]]})
    locked = out_dir / "locked.xlsx"
    ok = out_dir / "ok.xlsx"
    _wb(locked, {"Sheet1": []})
    _wb(ok, {"Sheet1": []})

    os.chmod(locked, 0o000)
    try:
        config = _cfg(str(master_dir), [
            _workbook("locked", "locked.xlsx", str(out_dir), [("Sheet1", "Sheet1")]),
            _workbook("ok", "ok.xlsx", str(out_dir), [("Sheet1", "Sheet1")]),
        ])
        p = WorkbookProcessor()
        cap = _Cap(p)
        p.configure(config, "master.xlsx", "Q1", ["locked", "ok"])
        p.run()
    finally:
        os.chmod(locked, 0o644)

    assert cap.errors == []
    statuses = {r.workbook_id: r.status for r in cap.finished}
    assert statuses["locked"] == "error"
    assert statuses["ok"] == "success"


@pytest.mark.skipif(sys.platform == "win32", reason="chmod-based permission simulation not reliable on Windows")
def test_unwritable_output_folder_emits_error_and_continues(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    locked_dir = tmp_path / "locked_out"
    locked_dir.mkdir()
    ok_dir = tmp_path / "ok_out"
    ok_dir.mkdir()

    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]]})
    _wb(locked_dir / "report.xlsx", {"Sheet1": []})
    _wb(ok_dir / "report.xlsx", {"Sheet1": []})

    os.chmod(locked_dir, 0o555)  # read+execute only, not writable
    try:
        config = _cfg(str(master_dir), [
            _workbook("locked", "report.xlsx", str(locked_dir), [("Sheet1", "Sheet1")]),
            _workbook("ok", "report.xlsx", str(ok_dir), [("Sheet1", "Sheet1")]),
        ])
        p = WorkbookProcessor()
        cap = _Cap(p)
        p.configure(config, "master.xlsx", "Q1", ["locked", "ok"])
        p.run()
    finally:
        os.chmod(locked_dir, 0o755)

    assert cap.errors == []
    statuses = {r.workbook_id: r.status for r in cap.finished}
    assert statuses["locked"] == "error"
    assert statuses["ok"] == "success"


# ── code-review additions ─────────────────────────────────────────────────────

def _wb_with_wrong_dimensions(path: Path, sheets: dict) -> None:
    """Create an xlsx whose dimension XML tag says only A1, forcing max_column/max_row to 1."""
    _wb(path, sheets)
    with zipfile.ZipFile(path, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}
    for key in list(entries):
        if re.match(r"xl/worksheets/sheet\d+\.xml", key):
            xml = entries[key].decode("utf-8")
            xml = re.sub(r"<dimension[^/]*/>", '<dimension ref="A1"/>', xml)
            entries[key] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)


def test_invalid_suffix_raises_value_error(tmp_path):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    config = _cfg(str(master_dir), [])
    p = WorkbookProcessor()
    with pytest.raises(ValueError, match="Suffix"):
        p.configure(config, "master.xlsx", "../../bad", [])


def test_unknown_workbook_id_in_selection_emits_error_and_continues(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    _wb(master_dir / "master.xlsx", {"Sheet1": [["val"]]})
    _wb(out_dir / "report.xlsx", {"Sheet1": []})

    config = _cfg(str(master_dir), [
        _workbook("known", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")]),
    ])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["unknown_id", "known"])
    p.run()

    assert cap.errors == []
    assert len(cap.finished) == 2
    statuses = {r.workbook_id: r.status for r in cap.finished}
    assert statuses["unknown_id"] == "error"
    assert statuses["known"] == "success"


def test_bad_dimension_metadata_copies_all_columns_and_clears_stale_data(tmp_path, qapp):
    master_dir = tmp_path / "master"
    master_dir.mkdir()
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    # Master has 3 columns but xlsx metadata says only A1
    _wb_with_wrong_dimensions(
        master_dir / "master.xlsx",
        {"Sheet1": [["a", "b", "c"], ["1", "2", "3"]]},
    )
    # Target has 4 columns; first 3 should be overwritten, 4th preserved
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Sheet1"
    ws.append(["old1", "old2", "old3", "keep"])
    ws.append(["x", "y", "z", "keep2"])
    ws.append(["stale", None, None, "keep3"])  # extra row that should be cleared
    book.save(out_dir / "report.xlsx")

    config = _cfg(str(master_dir), [_workbook("wb1", "report.xlsx", str(out_dir), [("Sheet1", "Sheet1")])])
    p = WorkbookProcessor()
    cap = _Cap(p)
    p.configure(config, "master.xlsx", "Q1", ["wb1"])
    p.run()

    assert cap.errors == []
    assert cap.finished[0].status == "success"

    out = openpyxl.load_workbook(out_dir / "report_Q1.xlsx")
    ws = out["Sheet1"]
    assert ws.cell(1, 1).value == "a"
    assert ws.cell(1, 2).value == "b"
    assert ws.cell(1, 3).value == "c"
    assert ws.cell(2, 1).value == "1"
    assert ws.cell(2, 2).value == "2"
    assert ws.cell(2, 3).value == "3"
    assert ws.cell(1, 4).value == "keep"    # outside paste zone — preserved
    assert ws.cell(2, 4).value == "keep2"
    assert ws.cell(3, 1).value is None      # stale row cleared
    assert ws.cell(3, 4).value == "keep3"   # outside paste zone — preserved
    out.close()
